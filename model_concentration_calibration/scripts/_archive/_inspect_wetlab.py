import openpyxl, os
files = [
  r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/rfp_tig0-7_extracted_SOURCE_CORRECTED_33h_clear.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/tig4567_extracted_data.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/pmlra_tig123_extracted_FOR_REVIEW.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/pmcy酶标仪数据提取.xlsx',
]
for f in files:
    print('=' * 90)
    print('FILE:', os.path.basename(f), os.path.getsize(f)//1024, 'KB')
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        print('  sheets:', wb.sheetnames[:6])
        for ws in wb.worksheets[:2]:
            rows = list(ws.iter_rows(min_row=1, max_row=8, values_only=True))
            print(f'  -- sheet {ws.title} ({ws.max_row} rows x {ws.max_column} cols)')
            for row in rows:
                print('    ', [str(c)[:22] if c is not None else '' for c in row])
    except Exception as e:
        print('  ERR', e)