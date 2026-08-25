import openpyxl
f = r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/rfp_tig0-7_extracted_SOURCE_CORRECTED_33h_clear.xlsx'
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
for sh in ['0123_0h_GFP', '0123_0h_RFP', '4567_6h_GFP', '4567_6h_RFP']:
    ws = wb[sh]
    print('='*70)
    print('SHEET', sh, 'dims', ws.max_row, 'x', ws.max_column)
    rows = list(ws.iter_rows(min_row=1, max_row=11, values_only=True))
    hdr = rows[0] if rows else []
    print('  header:', [str(c) if c is not None else '' for c in hdr])
    for row in rows[1:6]:
        print('  ', [str(c)[:14] if c is not None else '' for c in row])
    # also last rows
    lst = list(ws.iter_rows(min_row=max(2, ws.max_row-2), max_row=ws.max_row, values_only=True))
    for row in lst: print('  last:', [str(c)[:14] if c is not None else '' for c in row])