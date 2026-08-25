import openpyxl, os
paths = [
  r'D:/AAAProject/IGEM/IGEM-dry/model_redo/data/external_raw/USGS_Large_Rivers_2017_Microcystin/qPCR_Standard_Curve_Information.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/model_redo/data/external_raw/USGS_Large_Rivers_2018_Microcystin/2018_qPCR_Standard_Curve_Information.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/model_redo/data/external_raw/USGS_Large_Rivers_2019_Microcystin/2019_qPCR_Standard_Curve_Information.xlsx',
  r'D:/AAAProject/IGEM/IGEM-dry/model_redo/data/external_raw/USGS_Large_Rivers_2018_Microcystin/2018_Cyanotoxins_Chl_Genetics_Data.xlsx',
]
for p in paths:
    print('===', os.path.basename(p), os.path.getsize(p)//1024, 'KB')
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets[:2]:
            rows = list(ws.iter_rows(max_row=8, values_only=True))
            print('  sheet:', ws.title, 'dims:', ws.max_row, 'x', ws.max_column)
            for row in rows[:6]:
                print('   ', [str(c)[:26] for c in row if c is not None][:8])
    except Exception as e:
        print('  ERR', e)