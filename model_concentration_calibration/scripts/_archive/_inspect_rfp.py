import openpyxl
f = r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/rfp_tig0-7_extracted_SOURCE_CORRECTED_33h_clear.xlsx'
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
print('SHEETS (%d):' % len(wb.sheetnames))
for s in wb.sheetnames: print('  ', s)
print()
ws = wb['ReadMe']
for row in ws.iter_rows(values_only=True):
    print('ReadMe:', [str(c)[:120] if c is not None else '' for c in row])
print()
ws = wb['Correction_Notes']
for row in ws.iter_rows(max_row=12, values_only=True):
    print('Notes:', [str(c)[:100] if c is not None else '' for c in row])