# loader.py — real wet-lab data ingestion (rfp_tig0-7 corrected workbook) + QC
import os
import numpy as np
import openpyxl

CONCS = np.array([0.0, 0.05, 0.1, 0.25, 0.5, 1.0, 2.0, 5.0])
TIMES = [0.0, 2.0, 3.0, 4.0, 5.0, 6.0, 9.0, 24.0, 33.0]
T_KEYS = {0.0: '0h', 2.0: '2h', 3.0: '3h', 4.0: '4h', 5.0: '5h', 6.0: '6h', 9.0: '9h', 24.0: '24h', 33.0: '33h'}
GROUPS = {'0123': [0, 1, 2, 3], '4567': [4, 5, 6, 7]}
DEFAULT_WB = r'D:/AAAProject/IGEM/IGEM-dry/docs/Introduction/rfp_tig0-7_extracted_SOURCE_CORRECTED_33h_clear.xlsx'


def _read_sheet(ws):
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        conc = float(row[0])
        vals = [float(v) for v in row[2:14]]
        rows[conc] = vals
    return rows


def load(wb_path=DEFAULT_WB):
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    recs = []
    for gname, triggers in GROUPS.items():
        for th in TIMES:
            tk = T_KEYS[th]
            g_sh = gname + '_' + tk + '_GFP'
            r_sh = gname + '_' + tk + '_RFP'
            if g_sh not in wb.sheetnames or r_sh not in wb.sheetnames:
                continue
            G = _read_sheet(wb[g_sh])
            R = _read_sheet(wb[r_sh])
            for ci, conc in enumerate(CONCS):
                if conc not in G or conc not in R:
                    continue
                for tr_i, tid in enumerate(triggers):
                    for rep in range(3):
                        col = tr_i * 3 + rep
                        g = G[conc][col]; r = R[conc][col]
                        recs.append((gname, th, float(conc), tid, rep + 1, g, r, g / r if r > 0 else np.nan))
    wb.close()
    arr = np.array(recs, dtype=float)
    return dict(group=np.array([g for g, *_ in recs]), t=arr[:, 1], C=arr[:, 2],
                trigger=arr[:, 3].astype(int), rep=arr[:, 4].astype(int),
                G=arr[:, 5], R=arr[:, 6], Q=arr[:, 7])


def fold_by_time(d):
    # fold(c,t,r) = Q(c,t,r)/Q0(t,r)  where Q0 = same group/time/trigger/rep at C=0 (row A)
    key = {}
    out = []
    for i in range(len(d['C'])):
        k = (d['group'][i], d['t'][i], d['trigger'][i], d['rep'][i], 0.0)
        key.setdefault(k, []).append(i)
    q0 = {k: d['Q'][v[0]] for k, v in key.items()}
    for i in range(len(d['C'])):
        k = (d['group'][i], d['t'][i], d['trigger'][i], d['rep'][i], 0.0)
        q0v = q0[k]
        out.append((d['group'][i], d['t'][i], d['C'][i], d['trigger'][i], d['rep'][i], d['G'][i], d['R'][i], d['Q'][i], d['Q'][i] / q0v if q0v > 0 else np.nan))
    arr = np.array(out, dtype=float)
    return dict(group=np.array([o[0] for o in out]), t=arr[:, 1], C=arr[:, 2], trigger=arr[:, 3].astype(int),
                rep=arr[:, 4].astype(int), G=arr[:, 5], R=arr[:, 6], Q=arr[:, 7], fold=arr[:, 8])


def qc(d):
    out = {}
    for tid in range(8):
        m = d['trigger'] == tid
        cv_list = []
        for t in np.unique(d['t']):
            for rep in [1, 2, 3]:
                sel = m & (d['t'] == t) & (d['rep'] == rep)
                if int(sel.sum()) >= 4:
                    rvals = d['R'][sel]
                    cv_list.append(float(np.std(rvals) / max(np.mean(rvals), 1e-9)))
        out[tid] = (len(cv_list), float(np.median(cv_list)) if cv_list else float('nan'))
    return out